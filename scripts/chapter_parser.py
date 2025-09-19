"""Utilities for extracting structured content from a Word chapter document.

This script parses a `.docx` file that contains a single chapter.  It extracts:

* Chapter title
* Chapter objectives
* Chapter introduction
* Numbered sections (e.g., 1.1, 1.2) and subsections (e.g., 1.1.1)
* Content that follows the "Chapter Recap" heading

Tables that appear in the numbered sections are preserved as row/column text.
Content labelled as tips, callouts, or sidebars is skipped.
"""

from __future__ import annotations

import argparse
import json
import re
from collections import OrderedDict
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple, Union

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


EXCLUDED_PREFIXES = ("tip:", "callout:", "sidebar:")
SECTION_NUMBER_PATTERN = re.compile(r"^(?P<number>\d+(?:\.\d+)+)\s*(?P<title>.*)")


def normalize_text(text: str) -> str:
    """Collapse consecutive whitespace and trim the text."""
    return re.sub(r"\s+", " ", text).strip()


def is_excluded_text(text: str) -> bool:
    """Return True if the text should be ignored (tips, callouts, sidebars)."""
    lowered = text.strip().lower()
    return any(lowered.startswith(prefix) for prefix in EXCLUDED_PREFIXES)


def iter_block_items(parent: Union[Document, Table, Paragraph]) -> Iterable[Union[Paragraph, Table]]:
    """Yield paragraphs and tables in document order under *parent*."""
    if isinstance(parent, Document):
        parent_elm = parent.element.body
    else:
        parent_elm = parent._element  # type: ignore[attr-defined]

    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def table_to_rows(table: Table) -> List[List[str]]:
    """Convert a python-docx table into a 2D list of normalized text."""
    rows: List[List[str]] = []
    for row in table.rows:
        row_values = [normalize_text(cell.text) for cell in row.cells]
        if any(value for value in row_values):
            rows.append(row_values)
    return rows


def first_non_empty_cell(table_rows: List[List[str]]) -> Optional[str]:
    for row in table_rows:
        for value in row:
            if value:
                return value
    return None


def parse_numbered_heading(text: str) -> Optional[Tuple[str, str, int]]:
    """Return the section number, title, and depth if *text* is numbered."""
    match = SECTION_NUMBER_PATTERN.match(text)
    if not match:
        return None

    number = match.group("number")
    title = normalize_text(match.group("title"))
    depth = len(number.split("."))
    return number, title, depth


class Section:
    """Representation of a chapter section or subsection."""

    def __init__(self, identifier: str, title: str) -> None:
        self.identifier = identifier
        self.title = title
        self.content: List[Dict[str, object]] = []
        self.subsections: "OrderedDict[str, Section]" = OrderedDict()

    def add_paragraph(self, paragraph_text: str) -> None:
        text = normalize_text(paragraph_text)
        if text:
            self.content.append({"type": "paragraph", "text": text})

    def add_table(self, table_rows: List[List[str]]) -> None:
        if table_rows:
            self.content.append({"type": "table", "rows": table_rows})

    def to_dict(self) -> Dict[str, object]:
        return {
            "id": self.identifier,
            "title": self.title,
            "content": self.content,
            "subsections": OrderedDict(
                (identifier, subsection.to_dict())
                for identifier, subsection in self.subsections.items()
            ),
        }


class ChapterExtractor:
    """Parse and organize chapter content from a Word document."""

    def __init__(self, document: Document) -> None:
        self.document = document
        self.chapter_title: Optional[str] = None
        self.objectives: List[str] = []
        self.introduction: List[str] = []
        self.sections: "OrderedDict[str, Section]" = OrderedDict()
        self.chapter_recap: List[Dict[str, object]] = []

        self._current_section_stack: List[Tuple[int, Section]] = []
        self._gather_objectives = False
        self._gather_introduction = False
        self._gather_recap = False

    def parse(self) -> None:
        for block in iter_block_items(self.document):
            if isinstance(block, Paragraph):
                self._handle_paragraph(block)
            elif isinstance(block, Table):
                self._handle_table(block)

    # Internal helpers -------------------------------------------------

    def _handle_paragraph(self, paragraph: Paragraph) -> None:
        text = normalize_text(paragraph.text)
        if not text:
            return

        if self.chapter_title is None and not is_excluded_text(text):
            self.chapter_title = text
            return

        lowered = text.lower()
        if lowered == "chapter objectives":
            self._gather_objectives = True
            self._gather_introduction = False
            return
        if lowered == "introduction":
            self._gather_introduction = True
            self._gather_objectives = False
            return
        if lowered == "chapter recap":
            self._gather_recap = True
            self._gather_objectives = False
            self._gather_introduction = False
            self._current_section_stack.clear()
            return

        numbered = parse_numbered_heading(text)
        if numbered:
            self._start_section(*numbered)
            self._gather_objectives = False
            self._gather_introduction = False
            return

        if is_excluded_text(text):
            return

        if self._gather_recap:
            self.chapter_recap.append({"type": "paragraph", "text": text})
        elif self._gather_objectives:
            self.objectives.append(text)
        elif self._gather_introduction:
            self.introduction.append(text)
        elif self._current_section_stack:
            self._current_section_stack[-1][1].add_paragraph(text)

    def _handle_table(self, table: Table) -> None:
        table_rows = table_to_rows(table)
        if not table_rows:
            return

        first_value = first_non_empty_cell(table_rows)
        if first_value and is_excluded_text(first_value):
            return

        if self._gather_recap:
            self.chapter_recap.append({"type": "table", "rows": table_rows})
        elif self._current_section_stack:
            self._current_section_stack[-1][1].add_table(table_rows)

    def _start_section(self, identifier: str, title: str, depth: int) -> None:
        section = Section(identifier, title or identifier)

        while self._current_section_stack and self._current_section_stack[-1][0] >= depth:
            self._current_section_stack.pop()

        if self._current_section_stack:
            parent_section = self._current_section_stack[-1][1]
            parent_section.subsections[identifier] = section
        else:
            self.sections[identifier] = section

        self._current_section_stack.append((depth, section))

    # Public API -------------------------------------------------------

    def to_dict(self) -> Dict[str, object]:
        return {
            "chapter_title": self.chapter_title,
            "objectives": self.objectives,
            "introduction": self.introduction,
            "sections": OrderedDict(
                (identifier, section.to_dict())
                for identifier, section in self.sections.items()
            ),
            "chapter_recap": self.chapter_recap,
        }


def extract_chapter(path: Path) -> Dict[str, object]:
    document = Document(path)
    extractor = ChapterExtractor(document)
    extractor.parse()
    return extractor.to_dict()


def build_argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Extract structured content from a chapter Word document.",
    )
    parser.add_argument("docx_path", type=Path, help="Path to the .docx file")
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        help="Optional path to save the extracted JSON data.",
    )
    parser.add_argument(
        "--pretty",
        action="store_true",
        help="Pretty-print the output JSON with indentation.",
    )
    return parser


def main(argv: Optional[List[str]] = None) -> None:
    parser = build_argument_parser()
    args = parser.parse_args(argv)

    chapter_data = extract_chapter(args.docx_path)
    json_kwargs = {"ensure_ascii": False}
    if args.pretty:
        json_kwargs["indent"] = 2

    json_output = json.dumps(chapter_data, **json_kwargs)

    if args.output:
        args.output.write_text(json_output, encoding="utf-8")
    else:
        print(json_output)


if __name__ == "__main__":
    main()
